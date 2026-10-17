from datetime import datetime

import openpyxl
import pytest

from src.excel_manager import (
    ExcelManager,
    ExcelManagerError,
    categorizar_nota_evidencia,
    extraer_numero_elemento,
)
from src.modelos import CriterioBusqueda


@pytest.mark.parametrize(
    "nota,categoria_esperada",
    [
        (None, "sin_dato"),
        ("", "sin_dato"),
        ("FOTO Y VIDEO OK", "con_evidencia"),
        ("foto y video ok", "con_evidencia"),
        ("video ok", "con_evidencia"),
        ("PEDIR EVIDENCIAS", "pendiente"),
        ("PENDIENTE VIDEO", "pendiente"),
        ("BUSCAR NUMERO DE ELEMENTO", "pendiente"),
        ("ANULADA", "no_aplica"),
        ("N/A", "no_aplica"),
        ("dicrepancia posterior a la hora sr", "revisar_manual"),
    ],
)
def test_categorizar_nota_evidencia(nota, categoria_esperada):
    assert categorizar_nota_evidencia(nota) == categoria_esperada


def _crear_manager(ruta_excel) -> ExcelManager:
    return ExcelManager(
        ruta_archivo=str(ruta_excel),
        nombre_hoja="Asignacion- SEP",
        columna_codigo_largo="COD-INCIDENCIA",
        columna_codigo_corto="Incidencia",
        columna_evidencia_manual="EVIDENCIA",
        columna_estado_automatico="Estado_Automatico",
        columna_ruta_foto="Ruta_Evidencia",
        columna_observaciones="Observaciones",
        columna_numero_aviso="NUMERO AVISO",
        columna_fecha_deteccion="Fecha Detección",
        columna_fecha_sr="Fecha SR",
    )


@pytest.mark.parametrize(
    "observaciones,esperado",
    [
        (None, None),
        ("", None),
        ("se encontró el elemento fusible dañado", None),  # tipo sin número -> no cuenta
        ("apertura del elemento F3092 por falla", "F3092"),
        ("revisado elemento 30350 correctamente", "30350"),
        ("elemento interruptor fuera de servicio", None),
    ],
)
def test_extraer_numero_elemento(observaciones, esperado):
    assert extraer_numero_elemento(observaciones) == esperado


def test_abrir_falla_con_mensaje_claro_si_no_existe_el_archivo(tmp_path):
    manager = _crear_manager(tmp_path / "no_existe.xlsx")
    with pytest.raises(ExcelManagerError, match="No existe el archivo"):
        manager.abrir()


def test_abrir_falla_si_falta_la_hoja(excel_temporal):
    manager = ExcelManager(
        ruta_archivo=str(excel_temporal),
        nombre_hoja="HOJA_QUE_NO_EXISTE",
        columna_codigo_largo="COD-INCIDENCIA",
        columna_codigo_corto="Incidencia",
        columna_evidencia_manual="EVIDENCIA",
        columna_estado_automatico="Estado_Automatico",
        columna_ruta_foto="Ruta_Evidencia",
    )
    with pytest.raises(ExcelManagerError, match="no existe"):
        manager.abrir()


def test_crea_columnas_nuevas_sin_tocar_evidencia_original(excel_temporal):
    manager = _crear_manager(excel_temporal)
    manager.abrir()

    filas = manager.leer_filas()
    assert len(filas) == 4
    # La nota manual original se mantiene intacta y legible
    assert filas[0].nota_evidencia_manual == "PEDIR EVIDENCIAS"
    assert filas[0].categoria_nota_manual == "pendiente"


def test_buscar_fila_por_codigo_largo_y_corto(excel_temporal):
    manager = _crear_manager(excel_temporal)
    manager.abrir()

    por_largo = manager.buscar_fila_por_codigo("011682220001")
    por_corto = manager.buscar_fila_por_codigo("1168222")

    assert por_largo is not None
    assert por_corto is not None
    assert por_largo.fila_excel == por_corto.fila_excel


def test_buscar_fila_codigo_inexistente_devuelve_none(excel_temporal):
    manager = _crear_manager(excel_temporal)
    manager.abrir()
    assert manager.buscar_fila_por_codigo("000000000000") is None


def test_actualizar_estado_automatico_no_sobreescribe_columna_evidencia(excel_temporal):
    manager = _crear_manager(excel_temporal)
    manager.abrir()

    fila = manager.buscar_fila_por_codigo("1168222")
    manager.actualizar_estado_automatico(fila.fila_excel, estado="Descargada", ruta_foto="/ruta/falsa.jpg")
    manager.guardar()

    wb = openpyxl.load_workbook(excel_temporal)
    ws = wb["Asignacion- SEP"]

    encabezados = [c.value for c in ws[1]]
    assert "Estado_Automatico" in encabezados
    assert "Ruta_Evidencia" in encabezados

    idx_evidencia = encabezados.index("EVIDENCIA") + 1
    assert ws.cell(row=fila.fila_excel, column=idx_evidencia).value == "PEDIR EVIDENCIAS"  # intacta

    idx_estado = encabezados.index("Estado_Automatico") + 1
    assert ws.cell(row=fila.fila_excel, column=idx_estado).value == "Descargada"


def test_guardar_no_borra_otras_hojas_del_archivo(excel_temporal):
    manager = _crear_manager(excel_temporal)
    manager.abrir()
    manager.guardar()

    wb = openpyxl.load_workbook(excel_temporal)
    assert "OCTUBRE" in wb.sheetnames
    assert wb["OCTUBRE"]["A1"].value == "esta hoja no debe desaparecer"


# --- Cascada de búsqueda: código -> elemento -> aviso -> no encontrado ---
# Esto reproduce el criterio real que Henry usa a mano.

def test_cascada_encuentra_por_codigo_primero(excel_temporal):
    manager = _crear_manager(excel_temporal)
    manager.abrir()

    fila, criterio = manager.buscar_fila_por_identificadores(codigo="1168222")
    assert criterio == CriterioBusqueda.CODIGO
    assert fila.codigo_corto == "1168222"


def test_cascada_cae_a_numero_elemento_si_no_hay_codigo(excel_temporal):
    manager = _crear_manager(excel_temporal)
    manager.abrir()

    # "F3092" no es código de incidencia, pero sí es el elemento de la fila 1168222
    fila, criterio = manager.buscar_fila_por_identificadores(
        codigo="F3092", numero_elemento="F3092", numero_aviso="F3092"
    )
    assert criterio == CriterioBusqueda.NUMERO_ELEMENTO
    assert fila.codigo_corto == "1168222"


def test_cascada_cae_a_numero_aviso_si_no_hay_codigo_ni_elemento(excel_temporal):
    manager = _crear_manager(excel_temporal)
    manager.abrir()

    fila, criterio = manager.buscar_fila_por_identificadores(
        codigo="5604912", numero_elemento="5604912", numero_aviso="5604912"
    )
    assert criterio == CriterioBusqueda.NUMERO_AVISO
    assert fila.codigo_corto == "1166339"


def test_cascada_no_encontrado_si_ningun_criterio_matchea(excel_temporal):
    manager = _crear_manager(excel_temporal)
    manager.abrir()

    fila, criterio = manager.buscar_fila_por_identificadores(
        codigo="000000", numero_elemento="000000", numero_aviso="000000"
    )
    assert criterio == CriterioBusqueda.NO_ENCONTRADO
    assert fila is None


def test_leer_filas_incluye_fechas_para_validacion_horaria(excel_temporal):
    manager = _crear_manager(excel_temporal)
    manager.abrir()

    fila = manager.buscar_fila_por_codigo("1168222")
    assert fila.fecha_deteccion == datetime(2026, 10, 4, 12, 52)
    assert fila.fecha_sr == datetime(2026, 10, 4, 17, 17)


def test_columna_opcional_con_nombre_incorrecto_lanza_error_claro(excel_temporal):
    manager = ExcelManager(
        ruta_archivo=str(excel_temporal),
        nombre_hoja="Asignacion- SEP",
        columna_codigo_largo="COD-INCIDENCIA",
        columna_codigo_corto="Incidencia",
        columna_evidencia_manual="EVIDENCIA",
        columna_estado_automatico="Estado_Automatico",
        columna_ruta_foto="Ruta_Evidencia",
        columna_observaciones="ColumnaQueNoExiste",
    )
    with pytest.raises(ExcelManagerError, match="ColumnaQueNoExiste"):
        manager.abrir()


def test_columna_opcional_vacia_se_salta_sin_error(excel_temporal):
    manager = ExcelManager(
        ruta_archivo=str(excel_temporal),
        nombre_hoja="Asignacion- SEP",
        columna_codigo_largo="COD-INCIDENCIA",
        columna_codigo_corto="Incidencia",
        columna_evidencia_manual="EVIDENCIA",
        columna_estado_automatico="Estado_Automatico",
        columna_ruta_foto="Ruta_Evidencia",
        columna_observaciones="",
        columna_numero_aviso="",
        columna_fecha_deteccion="",
        columna_fecha_sr="",
    )
    manager.abrir()  # no debe lanzar nada
    filas = manager.leer_filas()
    assert all(f.numero_elemento is None and f.numero_aviso is None for f in filas)


def test_columnas_nuevas_van_tras_el_ultimo_encabezado_aunque_max_column_sea_enorme(excel_temporal):
    import openpyxl
    wb = openpyxl.load_workbook(excel_temporal)
    wb["Asignacion- SEP"].cell(row=1, column=16383).number_format = "General"  # infla max_column
    wb.save(excel_temporal)

    mgr = ExcelManager(str(excel_temporal), "Asignacion- SEP", "COD-INCIDENCIA", "Incidencia",
                       "EVIDENCIA", "Estado_Automatico", "Ruta_Evidencia")
    mgr.abrir()
    mgr.guardar()
    ws = openpyxl.load_workbook(excel_temporal)["Asignacion- SEP"]
    assert ws.cell(row=1, column=9).value == "Estado_Automatico"
    assert ws.cell(row=1, column=10).value == "Ruta_Evidencia"


def test_buscar_por_candidatos_respeta_la_cascada_entre_todos_los_numeros(excel_temporal):
    mgr = ExcelManager(str(excel_temporal), "Asignacion- SEP", "COD-INCIDENCIA", "Incidencia",
                       "EVIDENCIA", "Estado_Automatico", "Ruta_Evidencia")
    mgr.abrir()
    # aviso de la fila 3 escrito primero, pero el código de la fila 2 tiene prioridad
    fila, criterio, cual = mgr.buscar_fila_por_candidatos(["5604912", "1168222"])
    assert (fila.fila_excel, criterio, cual) == (2, CriterioBusqueda.CODIGO, "1168222")
    fila, criterio, cual = mgr.buscar_fila_por_candidatos(["9999999", "F3092"])
    assert (fila.fila_excel, criterio, cual) == (2, CriterioBusqueda.NUMERO_ELEMENTO, "F3092")
    fila, criterio, cual = mgr.buscar_fila_por_candidatos(["5604912"])
    assert (fila.fila_excel, criterio) == (3, CriterioBusqueda.NUMERO_AVISO)
    assert mgr.buscar_fila_por_candidatos(["0000000"]) == (None, CriterioBusqueda.NO_ENCONTRADO, None)


def _excel_con_provincias(tmp_path):
    import openpyxl
    ruta = tmp_path / "prov.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Asignacion- SEP"
    ws.append(["COD-INCIDENCIA", "Incidencia", "EVIDENCIA", "NUMERO AVISO", "Provincia"])
    ws.append(["011000010001", "1100001", None, 5600001, "VERAGUAS"])
    ws.append(["011000020001", "1100002", None, 5600002, "COCLÉ"])
    ws.append(["011000030001", "1100003", None, 5600003, "SIN DATOS"])
    wb.save(ruta)
    return ruta


def test_filtro_por_provincia_y_sin_datos(tmp_path):
    ruta = _excel_con_provincias(tmp_path)
    mgr = ExcelManager(str(ruta), "Asignacion- SEP", "COD-INCIDENCIA", "Incidencia", "EVIDENCIA",
                       "Estado_Automatico", "Ruta_Evidencia", columna_observaciones=None,
                       columna_numero_aviso="NUMERO AVISO", columna_fecha_deteccion=None,
                       columna_fecha_sr=None, columna_provincia="Provincia")
    mgr.abrir()
    assert mgr.resumen_por_provincia() == {"VERAGUAS": 1, "COCLE": 1, "SIN DATOS": 1}
    veraguas = {"VERAGUAS"}
    assert mgr.buscar_fila_por_candidatos(["1100001"], veraguas)[0].fila_excel == 2
    assert mgr.buscar_fila_por_candidatos(["1100002"], veraguas)[0] is None  # es de Coclé
    assert mgr.buscar_fila_por_candidatos(["1100002"], {"COCLE"})[0].fila_excel == 3  # tilde no importa
    assert mgr.buscar_fila_por_candidatos(["1100003"], veraguas)[0].fila_excel == 4  # SIN DATOS no se descarta
    assert mgr.buscar_fila_por_candidatos(["1100002"], None)[0].fila_excel == 3  # sin filtro


def test_provincias_de_grupo():
    from src.config import MatchingConfig, provincias_de_grupo

    class Cfg:  # basta con el atributo que usa la función
        matching = MatchingConfig("x", 15, provincias_por_grupo={
            "Veraguas": ["VERAGUAS"], "Coclé": ["Cocle"], "MECOPT(Oeste)": ["Panamá Oeste"]})
    assert provincias_de_grupo(Cfg, "CFFM_Veraguas_Sea_Central-Reporte de videos") == {"VERAGUAS"}
    assert provincias_de_grupo(Cfg, "CFFM_Coclé_Sea_Central") == {"COCLE"}
    assert provincias_de_grupo(Cfg, "Incidencias CFFM-MECOPT(Oeste)") == {"PANAMA OESTE"}
    assert provincias_de_grupo(Cfg, "Otro grupo") is None
