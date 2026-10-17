"""Tests del orquestador: _procesar_resultados y ejecutar_automatizacion con WhatsApp simulado."""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from src import main
from src.config import AppConfig, ExcelConfig, LoggingConfig, MatchingConfig, NotificacionesConfig, RutasConfig, WhatsAppConfig
from src.excel_manager import ExcelManager
from src.modelos import Mensaje, ResultadoMatch, TipoOrigenMatch

LOGGER = logging.getLogger("incidencias_wa")


def _cfg(tmp_path: Path, ruta_excel: Path, grupo: str = "Grupo Test") -> AppConfig:
    return AppConfig(
        whatsapp=WhatsAppConfig(grupos=[grupo], carpeta_sesion=str(tmp_path / "ses"), headless=True),
        excel=ExcelConfig(
            ruta_archivo=str(ruta_excel), nombre_hoja="Asignacion- SEP",
            columna_codigo_largo="COD-INCIDENCIA", columna_codigo_corto="Incidencia",
            columna_evidencia_manual="EVIDENCIA", columna_estado_automatico="Estado_Automatico",
            columna_ruta_foto="Ruta_Evidencia", columna_observaciones="Observaciones",
            columna_numero_aviso="NUMERO AVISO", columna_fecha_deteccion="Fecha Detección",
            columna_fecha_sr="Fecha SR",
        ),
        matching=MatchingConfig(patron_codigo_regex=r"\b\d{6,8}\b", ventana_minutos_mensaje_separado=15, margen_horas_evidencia=4),
        notificaciones=NotificacionesConfig(activas=False),
        rutas=RutasConfig(str(tmp_path / "evid"), "Incidencias_Identificadas", "Sin_Codigo_Asignado"),
        logging=LoggingConfig(str(tmp_path / "logs"), "INFO"),
    )


def _excel(cfg: AppConfig) -> ExcelManager:
    e = cfg.excel
    mgr = ExcelManager(
        e.ruta_archivo, e.nombre_hoja, e.columna_codigo_largo, e.columna_codigo_corto,
        e.columna_evidencia_manual, e.columna_estado_automatico, e.columna_ruta_foto,
        e.columna_observaciones, e.columna_numero_aviso, e.columna_fecha_deteccion, e.columna_fecha_sr,
    )
    mgr.abrir()
    return mgr


def _res(codigo, ts):
    msg = Mensaje(remitente="Juan", timestamp=ts, texto=f"inc {codigo}" if codigo else "", tiene_adjunto=True, tipo_adjunto="imagen")
    origen = TipoOrigenMatch.MISMO_MENSAJE if codigo else TipoOrigenMatch.SIN_CODIGO
    return ResultadoMatch(mensaje=msg, codigo=codigo, origen=origen)


def _estado(ruta: Path, fila: int):
    ws = openpyxl.load_workbook(ruta)["Asignacion- SEP"]
    cols = {c.value: c.column for c in ws[1]}
    return ws.cell(row=fila, column=cols["Estado_Automatico"]).value


def test_procesar_resultados_por_codigo_aviso_y_no_encontrado(excel_temporal, tmp_path, caplog):
    cfg = _cfg(tmp_path, excel_temporal)
    mgr = _excel(cfg)
    dentro = datetime(2026, 10, 4, 14, 0)  # entre Fecha Detección y Fecha SR de la fila 2
    resultados = [
        _res("1168222", dentro),                        # código corto -> fila 2
        _res("5604912", datetime(2026, 9, 24, 17, 0)),  # aviso -> fila 3
        _res("9999999", dentro),                        # no existe
        _res(None, dentro),                             # sin código
    ]
    with caplog.at_level("INFO", logger="incidencias_wa"):
        main._procesar_resultados(resultados, mgr, cfg, LOGGER)
    mgr.guardar()

    assert _estado(excel_temporal, 2) == "Detectada_pendiente_descarga"
    assert _estado(excel_temporal, 3) == "Detectada_pendiente_descarga"
    assert _estado(excel_temporal, 4) is None
    assert "NO ENCONTRADA" in caplog.text
    assert "criterio: numero_aviso" in caplog.text


def test_procesar_resultados_fuera_de_ventana_se_advierte(excel_temporal, tmp_path, caplog):
    cfg = _cfg(tmp_path, excel_temporal)
    mgr = _excel(cfg)
    with caplog.at_level("WARNING", logger="incidencias_wa"):
        main._procesar_resultados([_res("1168222", datetime(2026, 10, 10, 9, 0))], mgr, cfg, LOGGER)
    assert "fuera de la ventana" in caplog.text


def test_ejecutar_automatizacion_extremo_a_extremo_con_sesion_simulada(excel_temporal, tmp_path, monkeypatch):
    cfg = _cfg(tmp_path, excel_temporal)
    mensajes = [Mensaje("Juan", datetime(2026, 10, 4, 14, 0), "inc 1168222", True, "imagen")]

    class FakeSesion:
        def __init__(self, *a, **k):
            self.pagina = object()

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return None

        def abrir_chat(self, nombre):
            self.nombre = nombre

    monkeypatch.setattr(main, "load_config", lambda: cfg)
    monkeypatch.setattr(main, "SesionWhatsApp", FakeSesion)
    monkeypatch.setattr(main, "leer_mensajes_recientes", lambda pagina, **kw: mensajes)
    main.ejecutar_automatizacion()

    wb = openpyxl.load_workbook(excel_temporal)
    assert "OCTUBRE" in wb.sheetnames  # no se dañan otras hojas
    assert _estado(excel_temporal, 2) == "Detectada_pendiente_descarga"


def test_ejecutar_automatizacion_aborta_con_grupo_de_ejemplo(excel_temporal, tmp_path, monkeypatch):
    cfg = _cfg(tmp_path, excel_temporal, grupo="NOMBRE_EXACTO_DEL_GRUPO")

    def no_debe_llamarse(*a, **k):
        pytest.fail("no debe abrir WhatsApp")

    monkeypatch.setattr(main, "load_config", lambda: cfg)
    monkeypatch.setattr(main, "SesionWhatsApp", no_debe_llamarse)
    main.ejecutar_automatizacion()


def test_varios_grupos_un_fallo_no_tumba_a_los_demas(excel_temporal, tmp_path, monkeypatch):
    cfg = _cfg(tmp_path, excel_temporal)
    cfg.whatsapp.grupos = ["Malo", "Bueno"]
    abiertos = []

    class FakeSesion:
        def __init__(self, *a, **k):
            self.pagina = object()

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return None

        def abrir_chat(self, nombre):
            if nombre == "Malo":
                raise RuntimeError("no existe")
            abiertos.append(nombre)

    monkeypatch.setattr(main, "load_config", lambda: cfg)
    monkeypatch.setattr(main, "SesionWhatsApp", FakeSesion)
    monkeypatch.setattr(
        main, "leer_mensajes_recientes",
        lambda pagina, **kw: [Mensaje("Juan", datetime(2026, 10, 4, 14, 0), "inc 1168222", True, "imagen")],
    )
    main.ejecutar_automatizacion()
    assert abiertos == ["Bueno"]
    assert _estado(excel_temporal, 2) == "Detectada_pendiente_descarga"
