from datetime import datetime
from pathlib import Path

import openpyxl

from src.excel.incidencias import NOMBRE_BASE, archivos_nuevos, fusionar, incorporar_nuevas, leer_base, preparar_carpeta
from src.excel.unificar import unificar

ENC = ["ASIGNACION", "COD-INCIDENCIA", "zona", "Asignado a", "causa_agrupadora", "estado", "excepcion", "Incidencia", "Instalación", "Descripcion",
       "Fecha Detección", "Fecha SR", "Tipo de Causa", "Causa", "Estado Actual", "Observaciones", "Duración SR", "Provincia", "Distrito",
       "Corregimiento", "NUMERO AVISO", "EVIDENCIA"]


def _fila(inc, asign=datetime(2026, 10, 5), aviso=None, evid=None, obs="obs", quien="Henry", prov="VERAGUAS"):
    d = dict(ASIGNACION=asign, **{"COD-INCIDENCIA": f"0{inc}0001", "Asignado a": quien, "Incidencia": inc, "Fecha Detección": datetime(2026, 10, 4, 10),
             "Fecha SR": datetime(2026, 10, 4, 14), "Observaciones": obs, "Provincia": prov, "NUMERO AVISO": aviso, "EVIDENCIA": evid})
    return [d.get(c) for c in ENC]


def _libro(ruta: Path, hojas: dict[str, list[list]]):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        ws.append(ENC)
        for f in filas:
            ws.append(f)
    wb.save(ruta)


def test_unificar_junta_hojas_funde_reasignaciones_y_no_pierde_notas(tmp_path):
    origen = tmp_path / "orig.xlsx"
    _libro(origen, {
        "Asignacion- SEP": [_fila(1168222, datetime(2026, 10, 5), evid="NO APARECEN EN MOVILIDAD"), _fila(1168222, datetime(2026, 10, 12), aviso="5620125", evid="OK"),
                            _fila(1169000)],
        "OCTUBRE": [_fila(1169001, quien="Henry"), _fila(1169000, obs="más completa")],
    })
    salida = tmp_path / "base.xlsx"
    r = unificar(origen, salida, "Asignacion- SEP", ("OCTUBRE",))
    assert r["incidencias"] == 3                                             # 1168222 (2 asignaciones), 1169000 (2 hojas) y 1169001
    base = leer_base(salida)
    f = base["1168222"]
    assert f["N° de asignaciones"] == 2 and f["NUMERO AVISO"] == "5620125"
    assert "NO APARECEN EN MOVILIDAD" in f["EVIDENCIA"] and "OK" in f["EVIDENCIA"]         # las dos notas se conservan
    assert "05/10/2026" in f["Historial de asignaciones"] and "12/10/2026" in f["Historial de asignaciones"]
    assert "Asignacion- SEP" in base["1169000"]["Origen"] and "OCTUBRE" in base["1169000"]["Origen"]


def test_incorporar_agrega_nuevas_actualiza_existentes_y_mueve_el_archivo(tmp_path):
    carpeta = preparar_carpeta(tmp_path / "Incidencias")
    _libro(tmp_path / "orig.xlsx", {"Asignacion- SEP": [_fila(1168222, evid="primera nota"), _fila(1169000)]})
    unificar(tmp_path / "orig.xlsx", carpeta / NOMBRE_BASE, "Asignacion- SEP", ())
    _libro(carpeta / "nuevas" / "actualizada.xlsx", {"OCTUBRE": [_fila(1168222, evid="primera nota ampliada", aviso="5620125"), _fila(1170555)]})
    assert [f.name for f in archivos_nuevos(carpeta)] == ["actualizada.xlsx"]
    r = incorporar_nuevas(carpeta)
    assert r["agregadas"] == 1 and r["actualizadas"] == 1 and r["archivos"] == 1
    base = leer_base(carpeta / NOMBRE_BASE)
    assert set(base) == {"1168222", "1169000", "1170555"}
    assert base["1168222"]["EVIDENCIA"] == "primera nota ampliada"          # la nota ampliada reemplaza a la anterior (no se duplica)
    assert base["1168222"]["NUMERO AVISO"] == "5620125" and base["1168222"]["N° de asignaciones"] == 1   # misma asignación: no cuenta como otra
    assert not archivos_nuevos(carpeta) and list((carpeta / "procesadas").glob("*_actualizada.xlsx")) and list((carpeta / "respaldo").glob("INCIDENCIAS_BASE_*.xlsx"))


def test_incorporar_es_idempotente_y_reasignacion_posterior_suma_historial(tmp_path):
    carpeta = preparar_carpeta(tmp_path / "Incidencias")
    _libro(tmp_path / "orig.xlsx", {"Asignacion- SEP": [_fila(1168222, datetime(2026, 10, 5))]})
    unificar(tmp_path / "orig.xlsx", carpeta / NOMBRE_BASE, "Asignacion- SEP", ())
    nuevo = {"OCTUBRE": [_fila(1168222, datetime(2026, 10, 12), evid="NO APARECEN", aviso="5620125")]}
    _libro(carpeta / "nuevas" / "a.xlsx", nuevo)
    assert incorporar_nuevas(carpeta)["actualizadas"] == 1
    _libro(carpeta / "nuevas" / "a.xlsx", nuevo)                              # el mismo archivo otra vez
    r2 = incorporar_nuevas(carpeta)
    assert r2["actualizadas"] == 0 and r2["sin_cambios"] == 1
    f = leer_base(carpeta / NOMBRE_BASE)["1168222"]
    assert f["N° de asignaciones"] == 2 and f["Historial de asignaciones"].count("\n") == 1   # la reasignación quedó una sola vez


def test_archivo_sin_incidencias_se_deja_en_nuevas(tmp_path):
    carpeta = preparar_carpeta(tmp_path / "Incidencias")
    wb = openpyxl.Workbook()
    wb.active.append(["algo", "otro"])
    wb.save(carpeta / "nuevas" / "raro.xlsx")
    r = incorporar_nuevas(carpeta)
    assert r["archivos"] == 0 and r["omitidos"] == ["raro.xlsx"] and (carpeta / "nuevas" / "raro.xlsx").exists()


def test_fusionar_agrega_columnas_extra_sin_perder_nada():
    base = {"1": {"Incidencia": "1", "ASIGNACION": datetime(2026, 10, 5), "natalia": "x"}}
    c = fusionar(base, [("n.xlsx", 2, {"Incidencia": 1, "ASIGNACION": datetime(2026, 10, 5), "natalia": "y", "columna_nueva": "z"})])
    assert c["actualizadas"] == 1 and base["1"]["natalia"] == "y" and base["1"]["columna_nueva"] == "z"


def test_la_base_tiene_dos_hojas_nuevas_van_a_nuevas_y_se_pasan_al_historico(tmp_path):
    from openpyxl import load_workbook

    from src.excel.incidencias import consolidar_nuevas

    carpeta = preparar_carpeta(tmp_path / "Incidencias")
    _libro(tmp_path / "orig.xlsx", {"Asignacion- SEP": [_fila(1168222), _fila(1169000)]})
    unificar(tmp_path / "orig.xlsx", carpeta / NOMBRE_BASE, "Asignacion- SEP", ())
    assert load_workbook(carpeta / NOMBRE_BASE).sheetnames == ["Histórico", "Nuevas"]
    _libro(carpeta / "nuevas" / "n.xlsx", {"OCTUBRE": [_fila(1170555), _fila(1168222, evid="nota ampliada")]})
    incorporar_nuevas(carpeta)
    base = leer_base(carpeta / NOMBRE_BASE)
    assert base["1170555"]["_hoja"] == "Nuevas" and base["1168222"]["_hoja"] == "Histórico"      # lo nuevo a «Nuevas»; lo conocido se actualiza donde está
    wb = load_workbook(carpeta / NOMBRE_BASE)
    assert wb["Nuevas"].max_row == 2 and wb["Histórico"].max_row == 3
    assert consolidar_nuevas(carpeta) == 1
    assert load_workbook(carpeta / NOMBRE_BASE)["Nuevas"].max_row == 1 and leer_base(carpeta / NOMBRE_BASE)["1170555"]["_hoja"] == "Histórico"
    assert consolidar_nuevas(carpeta) == 0


def test_formato_filas_bajas_y_columnas_anchas_sin_ajustar_texto(tmp_path):
    from openpyxl import load_workbook

    _libro(tmp_path / "orig.xlsx", {"Asignacion- SEP": [_fila(1168222, obs="texto muy largo " * 30)]})
    unificar(tmp_path / "orig.xlsx", tmp_path / "b.xlsx", "Asignacion- SEP", ())
    ws = load_workbook(tmp_path / "b.xlsx")["Histórico"]
    hdr = [c.value for c in ws[1]]
    col = hdr.index("Observaciones") + 1
    assert ws.row_dimensions[2].height == 18 and not ws.cell(row=2, column=col).alignment.wrap_text
    assert ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width >= 100


def test_el_estado_del_cruce_se_escribe_dentro_de_la_base_en_las_dos_hojas(tmp_path):
    from types import SimpleNamespace

    from openpyxl import load_workbook

    from src.excel.reporte import escribir_estado_en_base

    carpeta = preparar_carpeta(tmp_path / "Incidencias")
    _libro(tmp_path / "orig.xlsx", {"Asignacion- SEP": [_fila(1168222)]})
    unificar(tmp_path / "orig.xlsx", carpeta / NOMBRE_BASE, "Asignacion- SEP", ())
    _libro(carpeta / "nuevas" / "n.xlsx", {"OCTUBRE": [_fila(1170555)]})
    incorporar_nuevas(carpeta)
    cfg = SimpleNamespace(excel=SimpleNamespace(columna_estado_automatico="Estado_Automatico", columna_ruta_foto="Ruta_Evidencia"))
    i = lambda corto, estado: {"corto": corto, "paquetes": 1, "estado": estado, "fotos": 2, "videos": 1, "ultima": None, "carpeta": f"C:/x/{corto}"}  # noqa: E731
    n = escribir_estado_en_base(cfg, [{"incidencias": [i("1168222", "Con evidencia"), i("1170555", "Sin evidencia: pedir")]}], carpeta / NOMBRE_BASE)
    assert n == 2
    wb = load_workbook(carpeta / NOMBRE_BASE)
    for hoja, esperado in (("Histórico", "Con evidencia"), ("Nuevas", "Sin evidencia: pedir")):
        ws = wb[hoja]
        idx = {c.value: c.column for c in ws[1]}
        assert ws.cell(row=2, column=idx["Estado_Automatico"]).value.startswith(esperado) and ws.cell(row=2, column=idx["Ruta_Evidencia"]).value.startswith("C:/x/")
    assert leer_base(carpeta / NOMBRE_BASE)["1168222"]["Estado_Automatico"].startswith("Con evidencia")   # y el estado sobrevive a incorporar/leer


def test_cargar_archivos_copia_incorpora_y_no_toca_el_original(tmp_path):
    from src.excel.incidencias import cargar_archivos

    carpeta = preparar_carpeta(tmp_path / "Incidencias")
    _libro(tmp_path / "orig.xlsx", {"Asignacion- SEP": [_fila(1168222)]})
    unificar(tmp_path / "orig.xlsx", carpeta / NOMBRE_BASE, "Asignacion- SEP", ())
    suelto = tmp_path / "Descargas" / "incidencias_nuevas.xlsx"
    suelto.parent.mkdir()
    _libro(suelto, {"OCTUBRE": [_fila(1170555), _fila(1170556)]})
    r = cargar_archivos(carpeta, [suelto, tmp_path / "no_existe.xlsx", tmp_path / "nota.txt"])
    assert r["agregadas"] == 2 and r["copiados"] == ["incidencias_nuevas.xlsx"] and set(r["invalidos"]) == {"no_existe.xlsx", "nota.txt"}
    assert suelto.exists()                                                  # el original sigue donde estaba
    assert leer_base(carpeta / NOMBRE_BASE)["1170555"]["_hoja"] == "Nuevas"
    solo = cargar_archivos(carpeta, [suelto], incorporar=False)             # con el bot corriendo: solo se deja el archivo
    assert solo["copiados"] and (carpeta / "nuevas" / solo["copiados"][0]).exists()


def test_cargar_otra_vez_el_mismo_contenido_con_otro_nombre_no_ensucia_el_historial(tmp_path):
    from src.excel.incidencias import cargar_archivos

    carpeta = preparar_carpeta(tmp_path / "Incidencias")
    filas = [_fila(1168222, aviso="5620125", evid="Nota única", obs="Texto de la observación"), _fila(1169000, obs="OTRA OBSERVACIÓN")]
    _libro(tmp_path / "orig.xlsx", {"Asignacion- SEP": filas})
    unificar(tmp_path / "orig.xlsx", carpeta / NOMBRE_BASE, "Asignacion- SEP", ())
    _libro(tmp_path / "con_otro_nombre.xlsx", {"OCTUBRE": [_fila(1168222, aviso="5620125", evid="nota única", obs="texto  de la observación"), _fila(1169000, obs="otra observación")]})
    r = cargar_archivos(carpeta, [tmp_path / "con_otro_nombre.xlsx"])
    assert r["agregadas"] == 0 and r["actualizadas"] == 0 and r["sin_cambios"] == 2         # solo mayúsculas/espacios: no es un cambio
    assert not leer_base(carpeta / NOMBRE_BASE)["1168222"].get("Historial de asignaciones")


def test_los_temporales_de_otros_programas_de_excel_se_ignoran(tmp_path):
    from src.excel.incidencias import es_temporal

    carpeta = preparar_carpeta(tmp_path / "Incidencias")
    _libro(carpeta / "nuevas" / "buena.xlsx", {"OCTUBRE": [_fila(1170555)]})
    for temporal in ("~$buena.xlsx", ".~lock.buena.xlsx#", "~buena.xlsx", "base.tmp.xlsx"):
        (carpeta / "nuevas" / temporal).write_bytes(b"lock")
    assert [f.name for f in archivos_nuevos(carpeta)] == ["buena.xlsx"]
    assert es_temporal("~$x.xlsx") and es_temporal(".~lock.x.xlsx#") and not es_temporal("incidencias_nuevas.xlsx")


def test_si_otro_programa_tiene_la_base_abierta_no_se_cae_y_deja_los_archivos(tmp_path, monkeypatch):
    from src.excel import unificar as U

    carpeta = preparar_carpeta(tmp_path / "Incidencias")
    _libro(tmp_path / "orig.xlsx", {"Asignacion- SEP": [_fila(1168222)]})
    unificar(tmp_path / "orig.xlsx", carpeta / NOMBRE_BASE, "Asignacion- SEP", ())
    _libro(carpeta / "nuevas" / "n.xlsx", {"OCTUBRE": [_fila(1170555)]})
    original = U._escribir
    monkeypatch.setattr(U, "_escribir", lambda *a, **k: (_ for _ in ()).throw(PermissionError("bloqueada por otro programa")))
    r = incorporar_nuevas(carpeta)
    assert r["base_bloqueada"] is True and (carpeta / "nuevas" / "n.xlsx").exists()          # no se movió: se reintenta después
    monkeypatch.setattr(U, "_escribir", original)
    r2 = incorporar_nuevas(carpeta)
    assert r2["agregadas"] == 1 and not r2["base_bloqueada"] and not archivos_nuevos(carpeta)


def test_asignacion_antigua_completa_columnas_vacias_y_recargar_no_repite(tmp_path):
    from src.excel.incidencias import cargar_archivos

    carpeta = preparar_carpeta(tmp_path / "Incidencias")
    _libro(tmp_path / "orig.xlsx", {"Asignacion- SEP": [_fila(1168222, datetime(2026, 10, 12), evid="nota actual")]})
    unificar(tmp_path / "orig.xlsx", carpeta / NOMBRE_BASE, "Asignacion- SEP", ())
    antigua = _fila(1168222, datetime(2026, 10, 5), evid="nota vieja distinta", obs="observación que la base no tenía")
    _libro(tmp_path / "a.xlsx", {"OCTUBRE": [antigua]})
    r1 = cargar_archivos(carpeta, [tmp_path / "a.xlsx"])
    assert r1["actualizadas"] == 1
    f = leer_base(carpeta / NOMBRE_BASE)["1168222"]
    assert f["ASIGNACION"] == datetime(2026, 10, 12) and "nota vieja distinta" in f["EVIDENCIA"]     # manda la actual; la antigua no se pierde
    _libro(tmp_path / "b.xlsx", {"OCTUBRE": [antigua]})                                            # el mismo contenido con otro nombre
    r2 = cargar_archivos(carpeta, [tmp_path / "b.xlsx"])
    assert r2["actualizadas"] == 0 and r2["sin_cambios"] == 1
