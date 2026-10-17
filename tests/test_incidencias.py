from datetime import datetime
from pathlib import Path

import openpyxl

from src.excel.incidencias import NOMBRE_BASE, archivos_nuevos, fusionar, incorporar_nuevas, leer_base, preparar_carpeta
from src.excel.unificar import unificar

ENC = ["ASIGNACION", "COD-INCIDENCIA", "zona", "Asignado a", "causa_agrupadora", "estado", "excepcion", "Incidencia", "Instalación", "Descripcion",
       "Fecha Detección", "Fecha SR", "Tipo de Causa", "Causa", "Estado Actual", "Observaciones", "Duración SR", "Provincia", "Distrito",
       "Corregimiento", "NUMERO AVISO", "EVIDENCIA"]


def _fila(inc, asign=datetime(2026, 10, 5), aviso=None, evid=None, obs="obs", quien="Henry", prov="VERAGUAS"):
    d = dict(ASIGNACION=asign, **{"COD-INCIDENCIA": f"0{inc}0001", "Asignado a": quien, "Incidencia": inc, "Fecha Detección": datetime(2026, 10, 4, 10),
             "Fecha SR": datetime(2026, 10, 4, 14), "Observaciones": obs, "Provincia": prov, "NUMERO AVISO": aviso, "EVIDENCIA": evid})
    return [d.get(c) for c in ENC]


def _libro(ruta: Path, hojas: dict[str, list[list]]):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        ws.append(ENC)
        for f in filas:
            ws.append(f)
    wb.save(ruta)


def test_unificar_junta_hojas_funde_reasignaciones_y_no_pierde_notas(tmp_path):
    origen = tmp_path / "orig.xlsx"
    _libro(origen, {
        "Asignacion- SEP": [_fila(1168222, datetime(2026, 10, 5), evid="NO APARECEN EN MOVILIDAD"), _fila(1168222, datetime(2026, 10, 12), aviso="5620125", evid="OK"),
                            _fila(1169000)],
        "OCTUBRE": [_fila(1169001, quien="Henry"), _fila(1169000, obs="más completa")],
    })
    salida = tmp_path / "base.xlsx"
    r = unificar(origen, salida, "Asignacion- SEP", ("OCTUBRE",))
    assert r["incidencias"] == 3                                             # 1168222 (2 asignaciones), 1169000 (2 hojas) y 1169001
    base = leer_base(salida)
    f = base["1168222"]
    assert f["N° de asignaciones"] == 2 and f["NUMERO AVISO"] == "5620125"
    assert "NO APARECEN EN MOVILIDAD" in f["EVIDENCIA"] and "OK" in f["EVIDENCIA"]         # las dos notas se conservan
    assert "05/10/2026" in f["Historial de asignaciones"] and "12/10/2026" in f["Historial de asignaciones"]
    assert "Asignacion- SEP" in base["1169000"]["Origen"] and "OCTUBRE" in base["1169000"]["Origen"]


def test_incorporar_agrega_nuevas_actualiza_existentes_y_mueve_el_archivo(tmp_path):
    carpeta = preparar_carpeta(tmp_path / "Incidencias")
    _libro(tmp_path / "orig.xlsx", {"Asignacion- SEP": [_fila(1168222, evid="primera nota"), _fila(1169000)]})
    unificar(tmp_path / "orig.xlsx", carpeta / NOMBRE_BASE, "Asignacion- SEP", ())
    _libro(carpeta / "nuevas" / "actualizada.xlsx", {"OCTUBRE": [_fila(1168222, evid="primera nota ampliada", aviso="5620125"), _fila(1170555)]})
    assert [f.name for f in archivos_nuevos(carpeta)] == ["actualizada.xlsx"]
    r = incorporar_nuevas(carpeta)
    assert r["agregadas"] == 1 and r["actualizadas"] == 1 and r["archivos"] == 1
    base = leer_base(carpeta / NOMBRE_BASE)
    assert set(base) == {"1168222", "1169000", "1170555"}
    assert base["1168222"]["EVIDENCIA"] == "primera nota ampliada"          # la nota ampliada reemplaza a la anterior (no se duplica)
    assert base["1168222"]["NUMERO AVISO"] == "5620125" and base["1168222"]["N° de asignaciones"] == 1   # misma asignación: no cuenta como otra
    assert not archivos_nuevos(carpeta) and list((carpeta / "procesadas").glob("*_actualizada.xlsx")) and list((carpeta / "respaldo").glob("INCIDENCIAS_BASE_*.xlsx"))


def test_incorporar_es_idempotente_y_reasignacion_posterior_suma_historial(tmp_path):
    carpeta = preparar_carpeta(tmp_path / "Incidencias")
    _libro(tmp_path / "orig.xlsx", {"Asignacion- SEP": [_fila(1168222, datetime(2026, 10, 5))]})
    unificar(tmp_path / "orig.xlsx", carpeta / NOMBRE_BASE, "Asignacion- SEP", ())
    nuevo = {"OCTUBRE": [_fila(1168222, datetime(2026, 10, 12), evid="NO APARECEN", aviso="5620125")]}
    _libro(carpeta / "nuevas" / "a.xlsx", nuevo)
    assert incorporar_nuevas(carpeta)["actualizadas"] == 1
    _libro(carpeta / "nuevas" / "a.xlsx", nuevo)                              # el mismo archivo otra vez
    r2 = incorporar_nuevas(carpeta)
    assert r2["actualizadas"] == 0 and r2["sin_cambios"] == 1
    f = leer_base(carpeta / NOMBRE_BASE)["1168222"]
    assert f["N° de asignaciones"] == 2 and f["Historial de asignaciones"].count("\n") == 1   # la reasignación quedó una sola vez


def test_archivo_sin_incidencias_se_deja_en_nuevas(tmp_path):
    carpeta = preparar_carpeta(tmp_path / "Incidencias")
    wb = openpyxl.Workbook()
    wb.active.append(["algo", "otro"])
    wb.save(carpeta / "nuevas" / "raro.xlsx")
    r = incorporar_nuevas(carpeta)
    assert r["archivos"] == 0 and r["omitidos"] == ["raro.xlsx"] and (carpeta / "nuevas" / "raro.xlsx").exists()


def test_fusionar_agrega_columnas_extra_sin_perder_nada():
    base = {"1": {"Incidencia": "1", "ASIGNACION": datetime(2026, 10, 5), "natalia": "x"}}
    c = fusionar(base, [("n.xlsx", 2, {"Incidencia": 1, "ASIGNACION": datetime(2026, 10, 5), "natalia": "y", "columna_nueva": "z"})])
    assert c["actualizadas"] == 1 and base["1"]["natalia"] == "y" and base["1"]["columna_nueva"] == "z"
