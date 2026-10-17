"""
Reporte en Excel (con formato) de lo que cruzó y de lo que se descargó. Hojas:
  Resumen        - tarjetas por grupo (asignadas, con evidencia, sin evidencia, fotos, videos) y gráfico por día.
  Por incidencia - una fila por incidencia asignada: estado, cuántos paquetes/fotos/videos, carpeta (hipervínculo).
  Por día        - cuánto llegó cada día, cuánto cruzó, cuánto se descartó y cuánto se bajó.
  Descargas      - cada intento de descarga (resultado y motivo del fallo).
  Sin cruce      - paquetes con foto/video que no cruzaron con ninguna incidencia (para revisar).
Solo LEE el Excel de incidencias y el índice; escribe un archivo nuevo, nunca el original.
"""
from __future__ import annotations

import csv
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import provincias_de_grupo
from .cruce import cruzar
from .descargas import etiqueta_identificador
from .matcher import asociar_mensajes

AZUL, GRIS = "1F3864", "F2F2F2"
VERDE, AMBAR, ROJO, NEUTRO = "C6EFCE", "FFEB9C", "FFC7CE", "E7E6E6"
EXT_VIDEO = {".mp4", ".mov", ".webm", ".mkv", ".3gp"}
_borde = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def _encabezado(ws, fila: int, titulos: list[str]) -> None:
    for c, t in enumerate(titulos, 1):
        celda = ws.cell(row=fila, column=c, value=t)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = _borde
    ws.row_dimensions[fila].height = 30


def _ajustar(ws, anchos: dict[int, float]) -> None:
    for col, ancho in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = ancho


def _tabla(ws, fila0: int, titulos: list[str], filas: list[list], anchos: dict[int, float]) -> None:
    _encabezado(ws, fila0, titulos)
    for i, datos in enumerate(filas, fila0 + 1):
        for c, v in enumerate(datos, 1):
            celda = ws.cell(row=i, column=c, value=v)
            celda.border = _borde
            celda.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
            if (i - fila0) % 2 == 0:
                celda.fill = PatternFill("solid", fgColor=GRIS)
    ws.freeze_panes = ws.cell(row=fila0 + 1, column=1)
    if filas:
        ws.auto_filter.ref = f"A{fila0}:{get_column_letter(len(titulos))}{fila0 + len(filas)}"
    _ajustar(ws, anchos)


def _color_estado(celda, estado: str) -> None:
    color = (VERDE if estado.startswith("Con evidencia") else AMBAR if estado.startswith("Revisar") else
             ROJO if estado.startswith("Sin evidencia") else NEUTRO)
    celda.fill = PatternFill("solid", fgColor=color)
    celda.font = Font(bold=True)


def _archivos(rutas: str | None) -> list[Path]:
    return [Path(r) for r in (rutas or "").split(";") if r]


def analizar_grupo(cfg, excel, indice, grupo: str, desde: datetime) -> dict:
    """Cruza un grupo y devuelve todo lo que necesitan las hojas del reporte."""
    m = cfg.matching
    provincias = provincias_de_grupo(cfg, grupo)
    mensajes = indice.mensajes(grupo)
    resultados = [r for r in asociar_mensajes(mensajes, m.patron_codigo_regex, m.ventana_minutos_mensaje_separado, m.patron_elemento_regex)
                  if r.mensaje.timestamp >= desde and r.mensaje.n_adjuntos > 0]
    cruces = cruzar(resultados, excel, m.margen_horas_evidencia, provincias)
    descargas = indice.estado_descarga(grupo)
    filas = excel._filtrar_por_provincias(excel.leer_filas(), provincias)

    por_fila: dict[int, list] = defaultdict(list)
    for c in cruces:
        if c.fila is not None:
            por_fila[c.fila.fila_excel].append(c)

    incidencias = []
    for f in filas:
        cs = por_fila.get(f.fila_excel, [])
        archivos = [a for c in cs for a in _archivos(descargas.get(c.match.mensaje.id_mensaje, (None, None))[1])]
        fotos = sum(a.suffix.lower() not in EXT_VIDEO for a in archivos)
        videos = len(archivos) - fotos
        if cs:
            estado = "Con evidencia" if any(c.estado == "ok" for c in cs) else "Revisar horario"
        elif f.fecha_deteccion and f.fecha_deteccion.date() < desde.date():
            estado = f"Fuera del rango consultado (antes del {desde:%d/%m})"
        else:
            estado = "Sin evidencia: pedir"
        if cs and not archivos:
            estado += " (sin descargar aún)"
        carpetas = sorted({a.parent for a in archivos})
        incidencias.append({
            "fila": f.fila_excel, "hoja": f.hoja or "Asignacion", "codigo": f.codigo_largo or f.codigo_corto or "", "corto": f.codigo_corto or "", "aviso": f.numero_aviso or "",
            "elemento": f.numero_elemento or "", "provincia": f.provincia or "", "deteccion": f.fecha_deteccion, "sr": f.fecha_sr,
            "estado": estado, "paquetes": len(cs), "fotos": fotos, "videos": videos,
            "primera": min((c.match.mensaje.timestamp for c in cs), default=None), "ultima": max((c.match.mensaje.timestamp for c in cs), default=None),
            "remitentes": ", ".join(sorted({c.match.mensaje.remitente for c in cs if c.match.mensaje.remitente})),
            "vias": ", ".join(sorted({c.match.origen.value for c in cs})), "carpeta": str(carpetas[0]) if carpetas else "",
            "nota": f.nota_evidencia_manual or "",
        })

    dias: dict[str, Counter] = defaultdict(Counter)
    for c in cruces:
        d = dias[f"{c.match.mensaje.timestamp:%Y-%m-%d}"]
        d["paquetes"] += 1
        d["adjuntos"] += c.match.mensaje.n_adjuntos
        d["cruzan" if c.fila is not None else "descartados"] += 1
        est = descargas.get(c.match.mensaje.id_mensaje, (None, None))
        if est[0] == "ok":
            arch = _archivos(est[1])
            d["fotos"] += sum(a.suffix.lower() not in EXT_VIDEO for a in arch)
            d["videos"] += sum(a.suffix.lower() in EXT_VIDEO for a in arch)
    sin_cruce = [c for c in cruces if c.fila is None]
    filas_descarga = []
    for c in cruces:
        ms = c.match.mensaje
        est, ruta = descargas.get(ms.id_mensaje, (None, None))
        if est is None and c.fila is None:
            continue  # lo que no cruza ni se bajó no es una "descarga"
        archivos = [a for a in _archivos(ruta) if a.exists()]
        tam = sum(a.stat().st_size for a in archivos) / 1e6
        resultado = ("pendiente (cruza, aún sin bajar)" if est is None else "error" if est == "error"
                     else "error" if not archivos else "ok")
        filas_descarga.append({
            "id": ms.id_mensaje, "dia": f"{ms.timestamp:%Y-%m-%d}", "hora": f"{ms.timestamp:%H:%M}", "remitente": ms.remitente or "",
            "identificador": etiqueta_identificador(c) or "(sin número)", "hoja": (c.fila.hoja or "Asignacion") if c.fila else "",
            "fila": c.fila.fila_excel if c.fila else "", "origen": c.match.origen.value, "esperados": ms.n_adjuntos, "bajados": len(archivos),
            "mb": round(tam, 2), "resultado": resultado, "carpeta": str(archivos[0].parent) if archivos else "",
        })
    return {"grupo": grupo, "incidencias": incidencias, "dias": dias, "sin_cruce": sin_cruce, "paquetes": len(cruces),
            "descargas": filas_descarga}


def _mb(v) -> float:
    try:
        return round(float(v or 0) / 1e6, 2)
    except (TypeError, ValueError):
        return 0.0


def _leer_csv(ruta: Path) -> list[dict]:
    if not ruta.exists():
        return []
    with open(ruta, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def generar_reporte(cfg, excel, indice, grupos: list[str], desde: datetime, ruta_salida: Path, carpeta_logs: Path | None = None,
                    analisis: list[dict] | None = None) -> Path:
    carpeta_logs = carpeta_logs or Path(cfg.logging.carpeta_logs)
    if analisis is None:
        analisis = [analizar_grupo(cfg, excel, indice, g, desde) for g in grupos]
    wb = Workbook()

    # ---- Resumen -------------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Reporte de evidencias WhatsApp - Incidencias"
    ws["A1"].font = Font(size=18, bold=True, color=AZUL)
    ws["A2"] = f"Generado: {datetime.now():%d/%m/%Y %H:%M}   |   Desde: {desde:%d/%m/%Y}   |   Excel: {Path(cfg.excel.ruta_archivo).name}"
    ws["A2"].font = Font(italic=True, color="595959")
    filas_res = []
    for a in analisis:
        inc = a["incidencias"]
        dentro = [i for i in inc if not i["estado"].startswith("Fuera")]
        con = sum(i["estado"].startswith(("Con evidencia", "Revisar")) for i in inc)
        filas_res.append([a["grupo"], len(inc), len(dentro), con, sum(i["estado"].startswith("Sin evidencia") for i in inc),
                          f"{con / len(dentro):.0%}" if dentro else "-", a["paquetes"], sum(i["fotos"] for i in inc), sum(i["videos"] for i in inc)])
    _tabla(ws, 4, ["Grupo", "Incidencias asignadas", "Dentro del rango", "Con evidencia", "Sin evidencia (pedir)", "% cubierto",
                   "Paquetes con foto/video", "Fotos descargadas", "Videos descargados"], filas_res,
           {1: 48, 2: 14, 3: 14, 4: 14, 5: 16, 6: 12, 7: 18, 8: 16, 9: 16})
    ws.freeze_panes = None
    ws.auto_filter.ref = None
    fila_g = 6 + len(filas_res)
    ws.cell(row=fila_g, column=1, value="Fotos y videos por día").font = Font(bold=True, size=12, color=AZUL)
    todos_dias = sorted({d for a in analisis for d in a["dias"]})
    _encabezado(ws, fila_g + 1, ["Día", "Cruzan con incidencia", "Descartados"])
    for i, d in enumerate(todos_dias, fila_g + 2):
        ws.cell(row=i, column=1, value=d)
        ws.cell(row=i, column=2, value=sum(a["dias"][d]["cruzan"] for a in analisis if d in a["dias"]))
        ws.cell(row=i, column=3, value=sum(a["dias"][d]["descartados"] for a in analisis if d in a["dias"]))
    if todos_dias:
        graf = BarChart()
        graf.type, graf.grouping, graf.overlap = "col", "stacked", 100
        graf.title, graf.height, graf.width = "Paquetes con foto/video por día", 8, 22
        graf.add_data(Reference(ws, min_col=2, max_col=3, min_row=fila_g + 1, max_row=fila_g + 1 + len(todos_dias)), titles_from_data=True)
        graf.set_categories(Reference(ws, min_col=1, min_row=fila_g + 2, max_row=fila_g + 1 + len(todos_dias)))
        ws.add_chart(graf, f"E{fila_g + 1}")

    # ---- Por incidencia -----------------------------------------------------------------------
    ws = wb.create_sheet("Por incidencia")
    filas = []
    for a in analisis:
        for i in a["incidencias"]:
            filas.append([a["grupo"][:40], i["hoja"], i["fila"], i["codigo"], i["corto"], i["aviso"], i["elemento"], i["provincia"],
                          i["deteccion"], i["sr"], i["estado"], i["paquetes"], i["fotos"], i["videos"], i["primera"], i["ultima"],
                          i["remitentes"], i["vias"], i["carpeta"], i["nota"]])
    _tabla(ws, 1, ["Grupo", "Hoja", "Fila Excel", "COD-INCIDENCIA", "Incidencia", "N° aviso", "Elemento", "Provincia", "Fecha detección", "Fecha SR",
                   "Estado", "Paquetes", "Fotos", "Videos", "Primera evidencia", "Última evidencia", "Remitentes", "Cómo se identificó",
                   "Carpeta", "Nota manual (Excel)"], filas,
           {1: 22, 2: 18, 3: 9, 4: 16, 5: 11, 6: 12, 7: 11, 8: 14, 9: 17, 10: 17, 11: 34, 12: 9, 13: 8, 14: 8, 15: 17, 16: 17, 17: 28, 18: 24, 19: 50, 20: 40})
    for r in range(2, len(filas) + 2):
        for col in (9, 10, 15, 16):
            ws.cell(row=r, column=col).number_format = "dd/mm/yyyy hh:mm"
        _color_estado(ws.cell(row=r, column=11), str(ws.cell(row=r, column=11).value))
        ruta = ws.cell(row=r, column=19).value
        if ruta:
            ws.cell(row=r, column=19).hyperlink = Path(ruta).resolve().as_uri()
            ws.cell(row=r, column=19).font = Font(color="0563C1", underline="single")

    # ---- Por día ------------------------------------------------------------------------------
    ws = wb.create_sheet("Por día")
    filas = [[a["grupo"][:40], d, c["paquetes"], c["adjuntos"], c["cruzan"], c["descartados"], c["fotos"], c["videos"]]
             for a in analisis for d, c in sorted(a["dias"].items())]
    _tabla(ws, 1, ["Grupo", "Día", "Paquetes con foto/video", "Fotos+videos (según chat)", "Cruzan con incidencia", "Descartados",
                   "Fotos descargadas", "Videos descargados"], filas, {1: 26, 2: 12, 3: 16, 4: 16, 5: 16, 6: 14, 7: 14, 8: 14})

    # ---- Descargas ----------------------------------------------------------------------------
    # Fuente: el índice (todo lo descargado, siempre completo). El CSV de detalle solo aporta el motivo de un fallo.
    ws = wb.create_sheet("Descargas")
    motivos = {d.get("id_mensaje"): d.get("error") for d in _leer_csv(carpeta_logs / "descargas_detalle.csv") if d.get("error")}
    filas = []
    for a in analisis:
        for d in sorted(a["descargas"], key=lambda x: (x["dia"], x["hora"])):
            filas.append([a["grupo"][:30], d["dia"], d["hora"], d["remitente"], d["identificador"], d["hoja"], d["fila"], d["origen"],
                          d["esperados"], d["bajados"], d["mb"], d["resultado"], motivos.get(d["id"], ""), d["carpeta"]])
    _tabla(ws, 1, ["Grupo", "Día", "Hora", "Remitente", "Identificador (aviso/incidencia)", "Hoja", "Fila", "Cómo se identificó", "Miniaturas en el chat",
                   "Bajados", "MB", "Resultado", "Motivo si falló", "Carpeta"], filas,
           {1: 20, 2: 11, 3: 7, 4: 18, 5: 34, 6: 14, 7: 7, 8: 18, 9: 10, 10: 9, 11: 8, 12: 28, 13: 40, 14: 60})
    for r in range(2, len(filas) + 2):
        res = str(ws.cell(row=r, column=12).value)
        ws.cell(row=r, column=12).fill = PatternFill("solid", fgColor=VERDE if res == "ok" else AMBAR if res.startswith(("parcial", "pendiente")) else ROJO)
        ruta = ws.cell(row=r, column=14).value
        if ruta:
            ws.cell(row=r, column=14).hyperlink = Path(ruta).resolve().as_uri()
            ws.cell(row=r, column=14).font = Font(color="0563C1", underline="single")

    # ---- Sin cruce ----------------------------------------------------------------------------
    ws = wb.create_sheet("Sin cruce")
    filas = []
    for a in analisis:
        for c in a["sin_cruce"]:
            ms = c.match.mensaje
            filas.append([a["grupo"][:30], ms.timestamp, ms.remitente, ms.n_adjuntos, c.estado, ", ".join(c.match.candidatos),
                          c.match.origen.value, (ms.texto or "")[:200]])
    _tabla(ws, 1, ["Grupo", "Fecha y hora", "Remitente", "Adjuntos", "Motivo", "Números en el chat", "Cómo se buscó", "Texto"],
           filas, {1: 22, 2: 17, 3: 22, 4: 9, 5: 16, 6: 30, 7: 20, 8: 60})
    for r in range(2, len(filas) + 2):
        ws.cell(row=r, column=2).number_format = "dd/mm/yyyy hh:mm"

    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta_salida)
    return ruta_salida


def _texto_estado(i: dict) -> str:
    base = i["estado"]
    if i["fotos"] or i["videos"]:
        base += f" | {i['fotos']} foto(s), {i['videos']} video(s)"
    if i["ultima"]:
        base += f" | última {i['ultima']:%d/%m %H:%M}"
    return base


def escribir_cruzado(cfg, analisis: list[dict], ruta_original: Path, ruta_salida: Path) -> int:
    """
    Copia del Excel de incidencias con el resultado del cruce escrito en las columnas Estado_Automatico y
    Ruta_Evidencia de CADA fila (hoja principal y hojas adicionales), coloreadas por estado. El original no se toca.
    Devuelve cuántas filas se escribieron. Los gráficos/imágenes del Excel original no se conservan en esta copia.
    """
    from openpyxl import load_workbook

    mejor: dict[tuple[str, int], dict] = {}
    for a in analisis:
        for i in a["incidencias"]:
            clave = (i["hoja"], i["fila"])
            previo = mejor.get(clave)
            # una fila "SIN DATOS" aparece en todos los grupos: gana el que sí encontró evidencia
            if previo is None or (i["paquetes"] > previo["paquetes"]):
                mejor[clave] = i
    wb = load_workbook(ruta_original)
    e = cfg.excel
    principal = wb[e.nombre_hoja]

    def columnas(ws, crear: bool):
        idx = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
        if crear:
            for nombre in (e.columna_estado_automatico, e.columna_ruta_foto):
                if nombre not in idx:
                    idx[nombre] = max(idx.values()) + 1
                    ws.cell(row=1, column=idx[nombre], value=nombre)
        return idx

    idx_main = columnas(principal, True)
    escritas = 0
    for hoja_nombre in (None, *e.hojas_extra):
        ws = principal if hoja_nombre is None else (wb[hoja_nombre] if hoja_nombre in wb.sheetnames else None)
        if ws is None:
            continue
        if hoja_nombre is None:
            idx = idx_main
        else:
            idx = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
            if e.columna_codigo_largo in idx:
                idx = columnas(ws, True)
            else:  # sin encabezado: mismas columnas que la hoja principal
                idx = idx_main
        clave_hoja = "Asignacion" if hoja_nombre is None else hoja_nombre
        for (hoja, fila), i in mejor.items():
            if (hoja or "Asignacion") != clave_hoja:
                continue
            celda = ws.cell(row=fila, column=idx[e.columna_estado_automatico], value=_texto_estado(i))
            _color_estado(celda, i["estado"])
            ws.cell(row=fila, column=idx[e.columna_ruta_foto], value=i["carpeta"] or None)
            escritas += 1
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    tmp = ruta_salida.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    import os
    os.replace(tmp, ruta_salida)
    return escritas
