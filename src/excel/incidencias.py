"""
Base de incidencias en una carpeta (data/Incidencias/):

    INCIDENCIAS_BASE.xlsx   la base: UNA hoja ("Incidencias"), una fila por incidencia, con historial de reasignaciones.
    nuevas/                 AQUÍ se ponen los Excel con incidencias nuevas (formato de asignación de siempre, o esta misma base).
    procesadas/             los archivos de `nuevas/` ya incorporados (con la fecha delante), por si hay que revisar algo.
    respaldo/               copia de la base antes de cada incorporación.
    original/               el Excel original de varias hojas, intacto.

`incorporar_nuevas` mezcla lo de `nuevas/` en la base sin perder nada: la incidencia nueva se agrega; si ya existía, manda la última
asignación y la anterior queda en «Historial de asignaciones» (se puede soltar el mismo archivo dos veces sin duplicar nada).
El vigilante lo ejecuta solo cuando detecta archivos nuevos; también se puede correr a mano (scripts/incorporar_incidencias.py).
"""
from __future__ import annotations

import logging
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from src.excel import unificar as U

logger = logging.getLogger("incidencias_wa")

NOMBRE_BASE = "INCIDENCIAS_BASE.xlsx"
SUBCARPETAS = ("nuevas", "procesadas", "respaldo", "original")
ENCABEZADO_ASIGNACION = [
    "ASIGNACION", "COD-INCIDENCIA", "zona", "Asignado a", "causa_agrupadora", "estado", "excepcion", "Incidencia", "Instalación", "Descripcion",
    "Fecha Detección", "Fecha SR", "Tipo de Causa", "Causa", "Estado Actual", "Observaciones", "Duración SR", "Provincia", "Distrito",
    "Corregimiento", "NUMERO AVISO", "EVIDENCIA", "WINDY", "IMHPA", "NOTICIA", "ANEXO A", "ANEXO C", "ANEXO B", "OT", "anexo c enviados", "zona (2)",
]
COLUMNAS_DEL_BOT = ("Estado_Automatico", "Ruta_Evidencia")


def carpeta_de(cfg) -> Path:
    """Carpeta de la base de incidencias según `excel.ruta_archivo` de config.yaml (relativa a la raíz del proyecto)."""
    from src.nucleo.rutas import R

    ruta = Path(cfg.excel.ruta_archivo)
    return (ruta if ruta.is_absolute() else R.raiz / ruta).resolve().parent


def preparar_carpeta(carpeta: Path) -> Path:
    for sub in SUBCARPETAS:
        (carpeta / sub).mkdir(parents=True, exist_ok=True)
    return carpeta


def archivos_nuevos(carpeta: Path) -> list[Path]:
    nuevas = carpeta / "nuevas"
    return sorted(f for f in nuevas.glob("*.xlsx") if not f.name.startswith("~$")) if nuevas.exists() else []


def _titulo_a_clave() -> dict[str, str]:
    return {titulo: clave for clave, titulo, _, _ in U.COLUMNAS}


def leer_base(ruta: Path) -> dict[str, dict]:
    """{incidencia: fila} con las columnas por su nombre interno."""
    if not ruta.exists():
        return {}
    wb = load_workbook(ruta, data_only=True, read_only=True)
    try:
        filas = [list(r) for r in wb[U.NOMBRE_HOJA].iter_rows(values_only=True)]
    finally:
        wb.close()  # en modo lectura openpyxl deja el archivo abierto: en Windows impediría moverlo o reemplazarlo
    mapa = _titulo_a_clave()
    claves = [mapa.get(str(t).strip(), str(t).strip()) if t is not None else "" for t in filas[0]]
    base: dict[str, dict] = {}
    for r in filas[1:]:
        d = {claves[j]: U._limpio(r[j]) for j in range(min(len(claves), len(r))) if claves[j] and not U._vacio(r[j])}
        inc = U._inc(d.get("Incidencia"))
        if inc:
            d["Incidencia"] = inc
            base[inc] = d
    return base


def leer_archivo_nuevo(ruta: Path) -> list[tuple[str, int, dict]]:
    """[(origen, fila, datos)] de un Excel nuevo: formato de asignación (con o sin encabezado) o la propia base unificada."""
    wb = load_workbook(ruta, data_only=True, read_only=True)
    try:
        return _filas_de_libro(wb, ruta)
    finally:
        wb.close()


def _filas_de_libro(wb, ruta: Path) -> list[tuple[str, int, dict]]:
    mapa = _titulo_a_clave()
    out: list[tuple[str, int, dict]] = []
    for hoja in wb.sheetnames:
        ws = wb[hoja]
        primera = next(ws.iter_rows(min_row=1, max_row=1, max_col=70, values_only=True), ())
        titulos = [str(v).strip() for v in primera if v is not None]
        if "Incidencia" in titulos and ("COD-INCIDENCIA" in titulos or "Origen (hojas del Excel original)" in titulos):
            filas = U._leer_hoja(ws, con_encabezado=True, max_col=70)
            for fila, d in filas:
                d = {mapa.get(k, k): v for k, v in d.items() if k not in COLUMNAS_DEL_BOT}
                if U._inc(d.get("Incidencia")):
                    out.append((ruta.name, fila, d))
        else:
            filas = U._leer_hoja(ws, con_encabezado=False, encabezado_ref=ENCABEZADO_ASIGNACION)
            validas = [(f, d) for f, d in filas if U._inc(d.get("Incidencia")) and str(d.get("COD-INCIDENCIA", "")).strip().isdigit()]
            if validas:
                logger.info("'%s' / hoja '%s': sin encabezado; se leyó con las columnas de la hoja de asignación.", ruta.name, hoja)
                out.extend((ruta.name, f, d) for f, d in validas)
    return out


def _combinar_notas(fila: dict, nueva: str | None, fecha) -> None:
    """Suma la nota nueva a la existente sin duplicar: si es la misma, nada; si es una versión más completa de la anterior, la reemplaza."""
    if not nueva:
        return
    nueva = str(nueva).strip()
    previa = str(fila.get("EVIDENCIA") or "").strip()
    if not previa:
        fila["EVIDENCIA"] = nueva
    elif nueva in previa:
        return
    elif previa in nueva:                                     # la nueva contiene a la anterior: es la misma nota ampliada
        fila["EVIDENCIA"] = nueva
    else:
        prefijo = f"{fecha:%d/%m}: " if isinstance(fecha, datetime) else ""
        fila["EVIDENCIA"] = f"{previa}\n{prefijo}{nueva}"


def fusionar(base: dict[str, dict], nuevas: list[tuple[str, int, dict]]) -> dict[str, int]:
    """Mezcla filas nuevas en `base` (se modifica). Devuelve cuántas incidencias se agregaron, actualizaron o ya estaban."""
    cuenta = {"agregadas": 0, "actualizadas": 0, "sin_cambios": 0}
    for origen, fila, d in nuevas:
        inc = U._inc(d.get("Incidencia"))
        marca = f"[{origen}, fila {fila}]"
        d = {k: v for k, v in d.items() if not U._vacio(v) and k not in COLUMNAS_DEL_BOT}
        if inc not in base:
            nueva = dict(d, Incidencia=inc, Origen=origen)
            avisos = ["".join(m) for m in U.RX_AVISO.findall(str(d.get("NUMERO AVISO", "")))]
            if len(avisos) > 1:
                nueva["NUMERO AVISO"], nueva["Otros avisos"] = avisos[-1], ", ".join(avisos[:-1])
            nueva["N° de asignaciones"] = 1
            base[inc] = nueva
            cuenta["agregadas"] += 1
            continue
        actual = base[inc]
        if marca in str(actual.get("Historial de asignaciones") or ""):
            cuenta["sin_cambios"] += 1                      # ese mismo archivo/fila ya se incorporó
            continue
        fecha_nueva, fecha_actual = d.get("ASIGNACION"), actual.get("ASIGNACION")
        es_mas_reciente = not (isinstance(fecha_nueva, datetime) and isinstance(fecha_actual, datetime)) or fecha_nueva >= fecha_actual
        diferencias = {k: v for k, v in d.items() if k not in ("Incidencia", "ASIGNACION", "Origen", "N° de asignaciones", "Historial de asignaciones")
                       and str(v).strip() != str(actual.get(k, "")).strip()}
        if not diferencias and fecha_nueva == fecha_actual:
            cuenta["sin_cambios"] += 1
            continue
        # historial: la fila anterior (si es la primera vez) y la nueva
        lineas = [actual["Historial de asignaciones"]] if actual.get("Historial de asignaciones") else []
        if not lineas and actual.get("ASIGNACION"):
            f0 = actual["ASIGNACION"]
            lineas.append(f"{f0:%d/%m/%Y}" + (f" [{actual.get('Origen', 'base')}]" if isinstance(f0, datetime) else " (sin fecha)"))
        partes = [f"{k}: {v}" for k, v in diferencias.items()]
        lineas.append((f"{fecha_nueva:%d/%m/%Y}" if isinstance(fecha_nueva, datetime) else "(sin fecha)") + f" {marca}" + (" · " + " | ".join(partes) if partes else ""))
        misma_asignacion = isinstance(fecha_nueva, datetime) and fecha_nueva == fecha_actual
        if misma_asignacion:
            lineas[-1] = lineas[-1].replace(f" {marca}", f" {marca} (actualización de la misma asignación)", 1)
        actual["Historial de asignaciones"] = "\n".join(lineas)
        if not misma_asignacion:
            actual["N° de asignaciones"] = int(actual.get("N° de asignaciones") or 1) + 1
        actual["Origen"] = ", ".join(dict.fromkeys([*str(actual.get("Origen", "")).split(", "), origen]))
        _combinar_notas(actual, d.get("EVIDENCIA"), fecha_nueva)
        nuevos_avisos = ["".join(m) for m in U.RX_AVISO.findall(str(d.get("NUMERO AVISO", "")))]
        if es_mas_reciente:
            for k, v in d.items():
                if k not in ("Incidencia", "EVIDENCIA", "NUMERO AVISO", "Origen", "N° de asignaciones", "Historial de asignaciones"):
                    actual[k] = v
            if nuevos_avisos:
                viejos = [a for a in ([actual.get("NUMERO AVISO")] + str(actual.get("Otros avisos") or "").split(", ")) if a and a not in nuevos_avisos]
                actual["NUMERO AVISO"] = nuevos_avisos[-1]
                actual["Otros avisos"] = ", ".join(dict.fromkeys([*nuevos_avisos[:-1], *viejos])) or None
        else:  # asignación anterior a la que ya hay: solo se suman sus avisos
            extra = [a for a in nuevos_avisos if a != str(actual.get("NUMERO AVISO", ""))]
            if extra:
                actual["Otros avisos"] = ", ".join(dict.fromkeys([*str(actual.get("Otros avisos") or "").split(", "), *extra])).strip(", ")
        cuenta["actualizadas"] += 1
    return cuenta


def incorporar_nuevas(carpeta: Path) -> dict:
    """Incorpora a la base todos los .xlsx de `nuevas/` y los mueve a `procesadas/`. Devuelve el resumen ({} si no había nada)."""
    carpeta = preparar_carpeta(carpeta)
    archivos = archivos_nuevos(carpeta)
    if not archivos:
        return {}
    ruta_base = carpeta / NOMBRE_BASE
    base = leer_base(ruta_base)
    if ruta_base.exists():
        shutil.copy2(ruta_base, carpeta / "respaldo" / f"INCIDENCIAS_BASE_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    total = {"archivos": 0, "agregadas": 0, "actualizadas": 0, "sin_cambios": 0, "omitidos": [], "sin_mover": []}
    procesar: list[Path] = []
    for f in archivos:
        try:
            filas = leer_archivo_nuevo(f)
        except Exception as e:  # noqa: BLE001 - un archivo roto no frena a los demás
            logger.error("No se pudo leer '%s': %s (se deja en nuevas/ para revisarlo)", f.name, e)
            total["omitidos"].append(f.name)
            continue
        if not filas:
            logger.warning("'%s' no trae incidencias reconocibles (¿falta la columna 'Incidencia'?); se deja en nuevas/.", f.name)
            total["omitidos"].append(f.name)
            continue
        c = fusionar(base, filas)
        logger.info("Incidencias nuevas '%s': %d filas -> %d agregadas, %d actualizadas, %d ya estaban.", f.name, len(filas), c["agregadas"], c["actualizadas"], c["sin_cambios"])
        for k in ("agregadas", "actualizadas", "sin_cambios"):
            total[k] += c[k]
        total["archivos"] += 1
        procesar.append(f)
    if procesar:
        U._escribir({k: v for k, v in base.items()}, ruta_base)
        total["sin_mover"] = []
        for f in procesar:
            try:
                f.replace(carpeta / "procesadas" / f"{datetime.now():%Y%m%d_%H%M%S}_{f.name}")
            except PermissionError:  # abierto en Excel: ya quedó incorporado (no se duplica); se moverá en la próxima corrida
                logger.warning("'%s' ya se incorporó, pero está abierto en Excel y no se pudo mover a procesadas/: ciérralo.", f.name)
                total["sin_mover"].append(f.name)
    return total
